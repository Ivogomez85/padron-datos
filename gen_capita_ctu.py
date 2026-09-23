#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gen_capita_ctu.py — genera el ARCHIVO MADRE de cápita asignada (CTU) a partir de los
Excel mensuales "Capitas_asignadas_CTU_YYYYMM.xlsx" que entrega PAMI.

Uso mensual:
    python3 gen_capita_ctu.py --padron padron_2026-08.csv.gz \
        --excel Capitas_asignadas_CTU_202607.xlsx Capitas_asignadas_CTU_202608.xlsx ... \
        --out capita_ctu.json [--previo capita_ctu.json]

- Cada Excel trae UN período (columna PERIODO). Si se pasa --previo, los meses que no
  vienen en los Excel nuevos se conservan del JSON anterior y los que sí vienen se
  REEMPLAZAN enteros (el Excel de un mes reextraído pisa al anterior).
- El padrón (csv.gz del repo padron-datos) sirve para traducir C_PRESTADOR -> N_SAP,
  N_LEGAJO, CUIT y descripción de módulo. Se usa el padrón más reciente disponible.

Reglas de lectura del Excel (validadas contra la cápita de referencia de jul/ago 2026):
- El grano real es prestador × módulo × AGENCIA ASIGNADA, identificada por
  (C_UGL_ASIGNADA, C_AGENCIA_ASIGNADA). La columna D_AGENCIA viene "cartesiana"
  (todas las agencias del país con ese código) y por eso se descarta: dentro de una
  misma agencia todas las filas repiten pct, afiliados y cápita.
- N_CAPITA_REF = PORCENTAJE_ASIGNADO × N_CANT_AGENCIA / 100 (cápita que aporta esa agencia).
- La cápita del módulo = suma de N_CAPITA_REF sobre las agencias DISTINTAS.
  (Arellano SAP 116041: una agencia 0002, 25% de 18.835 = 4.709 ✓;
   Clínica Integral 35928 módulo 543: 3.586 + 2.822 = 6.408 ✓)
- La cápita es un número de personas, no dinero, y entre módulos de un mismo bloque
  no se suma: se promedia (eso lo hace el motor del HTML).
"""
import argparse, csv, gzip, io, json, os, sys, datetime, collections, re
import openpyxl

RUBRO_RANGOS = [
    (lambda n: n in (36, 37, 38, 370, 380), "Oftalmología"),
    (lambda n: n in (2, 3, 22, 23, 24), "Imágenes"),
    (lambda n: n == 69 or 500 <= n <= 699, "Especialidades"),
]
def rubro(mod):
    try: n = int(mod)
    except: return "Otros"
    for f, r in RUBRO_RANGOS:
        if f(n): return r
    return "Otros"

import unicodedata
def _norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().upper()
    return " ".join(s.split())
LOC_UGL = collections.defaultdict(collections.Counter)   # localidad normalizada -> Counter(UGL): cuántas bocas del padrón hay ahí
STOP = {"DE", "DEL", "LA", "EL", "LAS", "LOS"}
def _clave(s): return " ".join(w for w in _norm(s).split() if w not in STOP)

def leer_padron(path):
    txt = gzip.open(path, "rb").read().decode("latin-1")
    rd = csv.DictReader(io.StringIO(txt), delimiter=";")
    mods, ugls, prest = {}, {}, {}
    for r in rd:
        _l = _clave(r.get("D_UBIC_GEO"))
        if _l: LOC_UGL[_l][r["c_ugl"].strip().zfill(2)] += 1
        m = r["C_MODULO_PAMI_"].strip()
        if m and m not in mods: mods[m] = r["D_MODULO_PAMI"].strip()
        u = r["c_ugl"].strip().zfill(2)
        if u not in ugls: ugls[u] = r["d_ugl"].strip()
        k = (r["C_PRESTADOR"].strip(), u)
        if k not in prest:
            prest[k] = {"sap": r["N_SAP"].strip(), "legajo": r["N_LEGAJO"].strip(),
                        "nombre": r["D_PRESTADOR"].strip(), "cuit": r["N_CUIT_CUIL"].strip()}
    return mods, ugls, prest

def leer_excel(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    H = {str(k).strip(): i for i, k in enumerate(rows[0]) if k is not None}
    req = ["PERIODO","C_UGL_PRESTADOR","C_PRESTADOR","D_PRESTADOR","C_MODULO_PAMI",
           "C_UGL_ASIGNADA","C_AGENCIA_ASIGNADA","PORCENTAJE_ASIGNADO","N_CANT_AGENCIA","N_CAPITA_REF"]
    falta = [c for c in req if c not in H]
    if falta: sys.exit("Excel %s: faltan columnas %s" % (path, falta))
    data = [r for r in rows[1:] if r[H["C_PRESTADOR"]] is not None]
    periodos = set(str(r[H["PERIODO"]]) for r in data)
    if len(periodos) != 1: sys.exit("Excel %s: más de un PERIODO (%s)" % (path, periodos))
    return periodos.pop(), H, data

AG_NOMBRES = collections.defaultdict(set)   # (ugl_asig, cod) -> nombres que trae D_AGENCIA (vienen cartesianos)

def procesar_mes(per, H, data, padron_prest, archivo):
    # prestador -> módulo -> agencia (ugl_asig, cod) -> fila
    P = collections.OrderedDict()
    for r in data:
        cp = str(r[H["C_PRESTADOR"]]).strip(); ugl = str(r[H["C_UGL_PRESTADOR"]]).strip().zfill(2)
        mod = str(r[H["C_MODULO_PAMI"]]).strip()
        ag = (str(r[H["C_UGL_ASIGNADA"]]).strip().zfill(2), str(r[H["C_AGENCIA_ASIGNADA"]]).strip())
        p = P.setdefault((cp, ugl), {"nombre": str(r[H["D_PRESTADOR"]]).strip(),
                                     "c_red": r[H["C_RED"]] if "C_RED" in H else None,
                                     "id_proceso": r[H["ID_PROCESO"]] if "ID_PROCESO" in H else None,
                                     "mods": collections.OrderedDict()})
        m = p["mods"].setdefault(mod, collections.OrderedDict())
        fila = {"ugl": ag[0], "cod": ag[1], "pct": r[H["PORCENTAJE_ASIGNADO"]],
                "afiliados": r[H["N_CANT_AGENCIA"]], "cap": r[H["N_CAPITA_REF"]]}
        if "D_AGENCIA" in H: AG_NOMBRES[ag].add(_norm(r[H["D_AGENCIA"]]))
        if ag in m:
            prev = m[ag]
            if (prev["pct"], prev["afiliados"], prev["cap"]) != (fila["pct"], fila["afiliados"], fila["cap"]):
                print("  AVISO %s %s mod %s agencia %s: valores distintos entre filas %s vs %s" % (per, cp, mod, ag, prev, fila), file=sys.stderr)
        else:
            m[ag] = fila
    out = {}
    sin_sap = []
    for (cp, ugl), p in P.items():
        pad = padron_prest.get((cp, ugl))
        if not pad: sin_sap.append((cp, ugl, p["nombre"]))
        modulos = []
        for mod in sorted(p["mods"], key=lambda x: int(x)):
            ags = list(p["mods"][mod].values())
            cap = sum((a["cap"] or 0) for a in ags)
            modulos.append({"mod": mod, "cap": int(round(cap)), "rubro": rubro(mod),
                            "n_agencias": len(ags),
                            "cruza_ugl": any(a["ugl"] != ugl for a in ags),
                            "agencias": ags})
        # promedio por bloque (misma regla que el motor del HTML: NO se suma entre módulos)
        acc = {}
        for m in modulos:
            acc.setdefault(m["rubro"], []).append(m["cap"])
        bloques = {k: int(round(sum(v) / len(v))) for k, v in acc.items()}
        out[(cp, ugl)] = {"nombre": p["nombre"], "c_red": p["c_red"], "id_proceso": p["id_proceso"],
                          "sap": pad["sap"] if pad else None, "legajo": pad["legajo"] if pad else None,
                          "cuit": pad["cuit"] if pad else None,
                          "mes": {"id_proceso": p["id_proceso"], "modulos": modulos, "bloques": bloques,
                                  "fuente": os.path.basename(archivo)}}
    return out, sin_sap

ROMAN = {i: r for i, r in enumerate(["I","II","III","IV","V","VI","VII","VIII","IX","X","XI","XII","XIII","XIV","XV","XVI","XVII","XVIII","XIX","XX","XXI","XXII","XXIII","XXIV","XXV","XXVI","XXVII","XXVIII","XXIX","XXX","XXXI","XXXII","XXXIII","XXXIV","XXXV","XXXVI","XXXVII","XXXVIII"], 1)}
PRE_AG = re.compile(r"^(CAP CABECERA|CAP|AGENCIA|BOCA DE ATENCION|CENTRO DE ATENCION PERSONALIZADA|OFICINA|DELEGACION)\s+")

def _ugls_de(nombre):
    if nombre.startswith("UGL "):
        m = re.match(r"UGL ([IVXL]+)", nombre)
        if m:
            for k, v in ROMAN.items():
                if v == m.group(1): return {str(k).zfill(2)}
        return set()
    base = _clave(PRE_AG.sub("", nombre))
    if not base: return set()
    cnt = collections.Counter()
    if base in LOC_UGL: cnt = collections.Counter(LOC_UGL[base])
    else:
        for l, us in LOC_UGL.items():
            if len(l) > 5 and (l in base or base in l): cnt.update(us)
    tot = sum(cnt.values())
    if not tot: return set()
    # solo UGLs donde la localidad es dominante (evita bocas sueltas de otra UGL en esa localidad)
    return set(u for u, n in cnt.items() if n >= 5 and n / tot >= 0.5)

def leer_maestro(path):
    """Maestro oficial de agencias (xlsx/csv) con columnas UGL, código de agencia y nombre."""
    filas = []
    if path.lower().endswith(".csv"):
        txt = open(path, encoding="utf-8-sig", errors="ignore").read()
        d = ";" if txt.split("\n", 1)[0].count(";") > txt.split("\n", 1)[0].count(",") else ","
        filas = list(csv.reader(io.StringIO(txt), delimiter=d))
    else:
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
        filas = [list(r) for r in ws.iter_rows(values_only=True)]
    H = {_norm(h): i for i, h in enumerate(filas[0]) if h is not None}
    def col(*names):
        for n in names:
            for k, i in H.items():
                if k == n or k.startswith(n): return i
        return None
    iu, ic, inm = col("C_UGL", "UGL"), col("C_AGENCIA", "CODIGO", "AGENCIA"), col("D_AGENCIA", "NOMBRE", "DESCRIPCION")
    if None in (iu, ic, inm): sys.exit("maestro de agencias: no encuentro columnas UGL / código / nombre en %s" % list(H))
    out = {}
    for r in filas[1:]:
        if r[iu] is None or r[ic] is None: continue
        u = str(r[iu]).strip().zfill(2); c = str(r[ic]).strip().zfill(4)
        out[(u, c)] = _norm(r[inm])
    return out

def resolver_agencias(maestro):
    """Devuelve {(ugl,cod): {"nombre","fuente"}}. D_AGENCIA del Excel viene cartesiano por código
    (trae todas las agencias del país con ese código), así que el nombre real se infiere cruzando
    la localidad del nombre con las localidades de la UGL en el padrón, con unicidad por código.
    Si hay maestro oficial, manda el maestro."""
    res = {}
    porcod = collections.defaultdict(list)
    for (u, c) in AG_NOMBRES: porcod[c].append(u)
    for cod, ugls in porcod.items():
        nombres = set()
        for u in ugls: nombres |= AG_NOMBRES[(u, cod)]
        cand = {n: _ugls_de(n) for n in nombres}
        pend, taken, asig = dict(cand), set(), {}
        changed = True
        while changed and pend:
            changed = False
            for n, us in list(pend.items()):
                us2 = [u for u in us if u in ugls and u not in taken]
                if len(us2) == 1:
                    asig[us2[0]] = (n, "unico" if len(us) == 1 else "inferido"); taken.add(us2[0]); del pend[n]; changed = True
            for u in ugls:
                if u in taken: continue
                cs = [n for n, us in pend.items() if u in us]
                if len(cs) == 1:
                    asig[u] = (cs[0], "inferido"); taken.add(u); del pend[cs[0]]; changed = True
        for u in ugls:
            if u in asig: res[(u, cod)] = {"nombre": asig[u][0], "fuente": asig[u][1]}
    for k, n in (maestro or {}).items():
        res[k] = {"nombre": n, "fuente": "maestro"}
    return res

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--padron", required=True, help="padron_YYYY-MM.csv.gz del repo padron-datos")
    ap.add_argument("--excel", nargs="+", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--previo", help="capita_ctu.json anterior: conserva los meses que no vengan en --excel")
    ap.add_argument("--agencias", help="maestro oficial de agencias (xlsx/csv: UGL, código, nombre); si se pasa, manda sobre lo inferido")
    a = ap.parse_args()

    mods_desc, ugls, padron_prest = leer_padron(a.padron)
    print("padrón: %d prestadores, %d módulos, %d UGL" % (len(padron_prest), len(mods_desc), len(ugls)))

    # prestadores: clave (cp, ugl) -> ficha + meses
    prest = collections.OrderedDict()
    meses_meta = collections.OrderedDict()
    if a.previo and os.path.exists(a.previo):
        prev = json.load(open(a.previo, encoding="utf-8"))
        meses_meta.update(prev.get("meses", {}))
        for p in prev.get("prestadores", []):
            prest[(p["c_prestador"], p["ugl"])] = p
        print("previo: %d prestadores, meses %s" % (len(prest), list(meses_meta)))

    for xl in a.excel:
        per, H, data = leer_excel(xl)
        res, sin_sap = procesar_mes(per, H, data, padron_prest, xl)
        # reemplazo entero del mes
        for p in prest.values():
            p["meses"].pop(per, None)
        for (cp, ugl), d in res.items():
            p = prest.get((cp, ugl))
            if not p:
                p = {"clave": None, "sap": d["sap"], "legajo": d["legajo"], "c_prestador": cp, "cuit": d["cuit"],
                     "ugl": ugl, "ugl_nombre": ugls.get(ugl, ""), "nombre": d["nombre"], "c_red": d["c_red"], "meses": {}}
                prest[(cp, ugl)] = p
            else:
                for k in ("sap", "legajo", "cuit"):
                    if not p.get(k) and d.get(k): p[k] = d[k]
                if d["nombre"]: p["nombre"] = d["nombre"]
            p["meses"][per] = d["mes"]
        for p in prest.values():
            p["clave"] = "%s||%s" % (p["sap"] or p["c_prestador"], p["ugl_nombre"])
        meses_meta[per] = {"archivo": os.path.basename(xl), "filas": len(data), "prestadores": len(res),
                           "modulos_prestador": sum(len(d["mes"]["modulos"]) for d in res.values()),
                           "cargado": datetime.date.today().isoformat()}
        print("%s: %d filas, %d prestadores, sin SAP en padrón: %d %s" % (per, len(data), len(res), len(sin_sap), sin_sap[:5]))

    # nombres de agencia: maestro (si hay) + inferencia; se escriben en cada fila de agencia
    maestro = leer_maestro(a.agencias) if a.agencias else {}
    prev_ag = {}
    if a.previo and os.path.exists(a.previo):
        try:
            for k, v in (json.load(open(a.previo, encoding="utf-8")).get("agencias") or {}).items():
                prev_ag[tuple(k.split("|"))] = v
        except Exception: pass
    agencias = resolver_agencias(maestro)
    for k, v in prev_ag.items():
        if k not in agencias or (v.get("fuente") == "maestro" and agencias[k]["fuente"] != "maestro"): agencias[k] = v
    for p in prest.values():
        for mm in p["meses"].values():
            for m in mm["modulos"]:
                for ag in m["agencias"]:
                    r = agencias.get((ag["ugl"], ag["cod"]))
                    ag["nombre"] = r["nombre"] if r else None
                    ag["fuente"] = r["fuente"] if r else None
    n_res = sum(1 for v in agencias.values()); n_tot = len(AG_NOMBRES)
    print("agencias: %d pares UGL+código, %d con nombre (%d maestro), %d sin resolver" % (n_tot, n_res, sum(1 for v in agencias.values() if v["fuente"] == "maestro"), n_tot - n_res))
    # sacar prestadores que quedaron sin ningún mes
    lista = [p for p in prest.values() if p["meses"]]
    lista.sort(key=lambda p: (int(p["ugl"]), p["nombre"]))
    meses_ord = collections.OrderedDict(sorted(meses_meta.items()))
    mods_usados = sorted(set(m["mod"] for p in lista for mm in p["meses"].values() for m in mm["modulos"]), key=int)
    doc = collections.OrderedDict([
        ("esquema", "capita-ctu/2"),
        ("generado", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("descripcion", "Cápita asignada (retribución CTU) por prestador, mes, módulo y agencia asignada. "
                        "Fuente: Excel mensual Capitas_asignadas_CTU de PAMI; SAP/legajo/CUIT resueltos contra el padrón. "
                        "La cápita es cantidad de personas (no dinero); entre módulos de un mismo bloque se promedia, no se suma."),
        ("clave", "sap||UGL_NOMBRE (misma llave que window.__CAPITA_EMBEBIDA). También se puede cruzar por c_prestador+ugl o legajo+ugl."),
        ("meses", meses_ord),
        ("ugl", ugls),
        ("agencias", collections.OrderedDict(("%s|%s" % k, v) for k, v in sorted(agencias.items()))),
        ("nota_agencias", "D_AGENCIA del Excel viene cartesiano por código; el nombre se infiere por localidad (fuente unico/inferido) o viene del maestro oficial (fuente maestro). 'afiliados' es el padrón de la agencia."),
        ("modulos", collections.OrderedDict((m, mods_desc.get(m, "")) for m in mods_usados)),
        ("prestadores", lista),
    ])
    json.dump(doc, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    print("OK -> %s (%d prestadores, %d KB)" % (a.out, len(lista), os.path.getsize(a.out) // 1024))

if __name__ == "__main__":
    main()
